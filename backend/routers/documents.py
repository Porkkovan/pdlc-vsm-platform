"""Project document upload for current-state assessment context."""
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, delete
import os
import uuid

from ..database.db import get_db
from ..database.models import ProjectDocument

router = APIRouter(prefix="/documents", tags=["documents"])

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "uploads")
ALLOWED_EXTENSIONS = {".pdf", ".xlsx", ".xls", ".csv", ".docx", ".doc", ".txt", ".json", ".pptx"}
CATEGORIES = [
    {"key": "general", "label": "General", "description": "General project documentation"},
    {"key": "process_doc", "label": "Process Documentation", "description": "SDLC process guides, runbooks, SOPs"},
    {"key": "metrics", "label": "Metrics & Reports", "description": "Performance reports, dashboards, KPI data"},
    {"key": "architecture", "label": "Architecture", "description": "Architecture diagrams, ADRs, tech stack docs"},
    {"key": "compliance", "label": "Compliance & Governance", "description": "Compliance frameworks, audit reports"},
    {"key": "tool_config", "label": "Tool Configuration", "description": "Tool setup docs, pipeline configs, CI/CD definitions"},
]


@router.get("/categories")
async def get_categories():
    return CATEGORIES


@router.get("/{project_id}")
async def list_documents(project_id: str, db: AsyncSession = Depends(get_db)):
    result = await db.execute(
        select(ProjectDocument)
        .where(ProjectDocument.project_id == project_id)
        .order_by(ProjectDocument.created_at.desc())
    )
    docs = result.scalars().all()
    return [_doc_dict(d) for d in docs]


@router.post("/{project_id}")
async def upload_document(
    project_id: str,
    file: UploadFile = File(...),
    category: str = Form("general"),
    db: AsyncSession = Depends(get_db),
):
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, f"File type {ext} not supported. Allowed: {', '.join(ALLOWED_EXTENSIONS)}")

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    file_id = str(uuid.uuid4())
    save_name = f"{file_id}{ext}"
    save_path = os.path.join(UPLOAD_DIR, save_name)

    content = await file.read()
    with open(save_path, "wb") as f:
        f.write(content)

    extracted = _extract_text(save_path, ext)

    doc = ProjectDocument(
        project_id=project_id,
        filename=file.filename,
        content_type=file.content_type,
        file_size=len(content),
        category=category,
        extracted_text=extracted,
        upload_path=save_path,
    )
    db.add(doc)
    await db.commit()
    await db.refresh(doc)
    return _doc_dict(doc)


@router.delete("/{project_id}/{doc_id}")
async def delete_document(project_id: str, doc_id: str, db: AsyncSession = Depends(get_db)):
    result = await db.execute(
        select(ProjectDocument).where(
            ProjectDocument.id == doc_id,
            ProjectDocument.project_id == project_id
        )
    )
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(404, "Document not found")
    if doc.upload_path and os.path.exists(doc.upload_path):
        try:
            os.remove(doc.upload_path)
        except OSError:
            pass
    await db.execute(
        delete(ProjectDocument).where(ProjectDocument.id == doc_id)
    )
    await db.commit()
    return {"deleted": True}


def _extract_text(path: str, ext: str) -> str:
    """Best-effort text extraction for agent consumption."""
    try:
        if ext == ".txt":
            with open(path, "r", errors="ignore") as f:
                return f.read()[:50000]
        if ext == ".csv":
            with open(path, "r", errors="ignore") as f:
                return f.read()[:50000]
        if ext == ".pdf":
            try:
                import PyPDF2
                with open(path, "rb") as f:
                    reader = PyPDF2.PdfReader(f)
                    text = "\n".join(page.extract_text() or "" for page in reader.pages[:50])
                return text[:50000]
            except ImportError:
                return "[PDF text extraction requires PyPDF2]"
        if ext in (".xlsx", ".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                lines = []
                for ws in wb.worksheets[:5]:
                    lines.append(f"--- Sheet: {ws.title} ---")
                    for row in ws.iter_rows(max_row=200, values_only=True):
                        lines.append("\t".join(str(c) if c is not None else "" for c in row))
                return "\n".join(lines)[:50000]
            except ImportError:
                return "[Excel extraction requires openpyxl]"
        if ext in (".docx",):
            try:
                import docx
                d = docx.Document(path)
                return "\n".join(p.text for p in d.paragraphs)[:50000]
            except ImportError:
                return "[DOCX extraction requires python-docx]"
    except Exception as e:
        return f"[Extraction error: {e}]"
    return ""


def _doc_dict(d: ProjectDocument) -> dict:
    return {
        "id": d.id,
        "project_id": d.project_id,
        "filename": d.filename,
        "content_type": d.content_type,
        "file_size": d.file_size,
        "category": d.category,
        "has_extracted_text": bool(d.extracted_text),
        "extracted_text_preview": (d.extracted_text or "")[:500],
        "created_at": d.created_at.isoformat() if d.created_at else None,
    }
